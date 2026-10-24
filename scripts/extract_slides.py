"""
extract_slides.py - Automated PowerPoint text and diagram OCR extractor.
Usage:
    python extract_slides.py <input_folder_or_file> <output_text_folder>
"""

import os
import sys
import io
from PIL import Image

try:
    from pptx import Presentation
except ImportError:
    print("Error: python-pptx is required. Install via: pip install python-pptx")
    sys.exit(1)

try:
    import winocr
    HAS_WINOCR = True
except ImportError:
    HAS_WINOCR = False

def convert_ppt_to_pptx_if_needed(file_path):
    """Converts legacy binary .ppt to .pptx using PowerPoint COM on Windows."""
    if not file_path.lower().endswith(".ppt"):
        return file_path

    pptx_path = file_path + "x"
    if os.path.exists(pptx_path):
        return pptx_path

    print(f"Converting legacy .ppt to .pptx: {os.path.basename(file_path)}")
    try:
        import win32com.client
        ppt_app = win32com.client.Dispatch("PowerPoint.Application")
        pres = ppt_app.Presentations.Open(os.path.abspath(file_path), WithWindow=False)
        pres.SaveAs(os.path.abspath(pptx_path), 24)  # 24 = ppSaveAsOpenXMLPresentation
        pres.Close()
        return pptx_path
    except Exception as e:
        print(f"Warning: COM conversion failed ({e}). Proceeding if possible.")
        return file_path

def ocr_image(image_bytes):
    """Runs OCR on image bytes using winocr or fallback."""
    if not HAS_WINOCR:
        return ""
    try:
        img = Image.open(io.BytesIO(image_bytes))
        res = winocr.recognize_pil_sync(img, "en")
        lines = [line["text"] for line in res.get("lines", [])]
        return "\n".join(lines).strip()
    except Exception:
        return ""

def process_presentation(pptx_path, out_txt_path):
    prs = Presentation(pptx_path)
    total_slides = len(prs.slides)
    print(f"Processing '{os.path.basename(pptx_path)}' ({total_slides} slides)...")

    full_text = []
    full_text.append(f"=== PRESENTATION: {os.path.basename(pptx_path)} ({total_slides} SLIDES) ===\n")

    for s_idx, slide in enumerate(prs.slides):
        slide_num = s_idx + 1
        slide_text = []
        slide_text.append(f"--- Slide {slide_num} ---")

        # Slide notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_text.append(f"[Presenter Notes]: {notes}")

        # Shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    slide_text.append(txt)

            elif shape.has_table:
                table_data = []
                for row in shape.table.rows:
                    row_txt = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    table_data.append(" | ".join(row_txt))
                if table_data:
                    slide_text.append("[Table Content]:\n" + "\n".join(table_data))

            elif shape.shape_type == 13:  # Picture
                try:
                    ocr_res = ocr_image(shape.image.blob)
                    if ocr_res:
                        slide_text.append(f"[Diagram OCR Text]:\n{ocr_res}")
                except Exception:
                    pass

        full_text.append("\n".join(slide_text) + "\n")

    with open(out_txt_path, "w", encoding="utf-8") as f:
        f.write("\n".join(full_text))
    print(f"Saved: {out_txt_path}")

def main():
    if len(sys.argv) < 3:
        print("Usage: python extract_slides.py <input_folder_or_file> <output_text_folder>")
        sys.exit(1)

    inp = sys.argv[1]
    out_dir = sys.argv[2]
    os.makedirs(out_dir, exist_ok=True)

    if os.path.isfile(inp):
        actual_path = convert_ppt_to_pptx_if_needed(inp)
        out_name = os.path.splitext(os.path.basename(inp))[0] + ".txt"
        process_presentation(actual_path, os.path.join(out_dir, out_name))
    elif os.path.isdir(inp):
        files = sorted(os.listdir(inp))
        for f in files:
            if f.lower().endswith((".pptx", ".ppt")):
                full_p = os.path.join(inp, f)
                actual_path = convert_ppt_to_pptx_if_needed(full_p)
                out_name = os.path.splitext(f)[0] + ".txt"
                process_presentation(actual_path, os.path.join(out_dir, out_name))

if __name__ == "__main__":
    main()
